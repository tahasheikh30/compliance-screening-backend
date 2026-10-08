"""Batch screening: reading Excel / CSV / Word files, background screening, privacy, cancel, and the downloads."""
import io
import threading
import time
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from app import batch_files
from tests import fixtures as fx
from tests.conftest import API_HEADERS, PENDING_HEADERS, USER2_HEADERS, USER_HEADERS

XLSX = "application/octet-stream"
HEADER = ["Full name", "Date of birth", "Nationality", "CNIC", "Father or husband", "Province"]


def make_xlsx(rows, header=HEADER) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str):
                c.data_type = "s"          # text stays text, even when it starts with "="
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def make_docx(rows, header=HEADER) -> bytes:
    def cell(t):
        return f"<w:tc><w:p><w:r><w:t>{t}</w:t></w:r></w:p></w:tc>"

    def tr(cells):
        return "<w:tr>" + "".join(cell(c) for c in cells) + "</w:tr>"

    xml = ('<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           "<w:body><w:p><w:r><w:t>Applicants</w:t></w:r></w:p><w:tbl>" + tr(header) + "".join(tr(r) for r in rows)
           + "</w:tbl></w:body></w:document>")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("word/document.xml", xml)
    return out.getvalue()


def post(client, data, name="applicants.xlsx", headers=USER_HEADERS, **params):
    qs = {"filename": name, **params}
    return client.post("/api/batch", params=qs, content=data, headers={**headers, "Content-Type": XLSX})


def wait(client, batch_id, headers=USER_HEADERS, timeout=20):
    end = time.time() + timeout
    while time.time() < end:
        body = client.get(f"/api/batches/{batch_id}", headers=headers).json()
        if body["status"] != "running":
            return body
        time.sleep(0.05)
    raise AssertionError("the batch did not finish")


ROWS = [
    ["Muhammad Ali Khan", "1975-03-04", "Pakistan", "42101-1234567-1", "", "Sindh"],   # on the UN list
    ["Completely Unrelated Person", "1990-01-01", "Pakistan", "", "", ""],
    ["", "1980-01-01", "", "", "", ""],                                                # no name
    ["Sara Noor", "", "", "", "", ""],
]


# ---- reading the file ------------------------------------------------------

def test_xlsx_batch_screens_every_row_and_keeps_bad_rows_visible(client):
    r = post(client, make_xlsx(ROWS))
    assert r.status_code == 202
    body = wait(client, r.json()["id"])
    assert body["status"] == "done" and body["total"] == 4 and body["done"] == 4
    rows = {x["row"]: x for x in body["rows"]}
    assert rows[2]["overall_status"] == "ESCALATE_TO_COMPLIANCE" and rows[2]["sanctions"] >= 1
    assert rows[3]["overall_status"] == "AUTO_CLEAR" and rows[3]["sanctions"] == 0
    assert rows[4]["state"] == "invalid" and rows[4]["overall_status"] is None and "Full name" in rows[4]["error"]
    assert rows[5]["state"] == "screened"
    assert body["counts"]["ESCALATE_TO_COMPLIANCE"] == 1 and body["counts"]["invalid"] == 1
    assert rows[2]["case_ref"].startswith("CS-")
    assert rows[2]["dob"] == "1975-03-04" and rows[2]["nationality"] == "Pakistan"      # for the case view


def test_each_row_is_an_ordinary_screening_in_history(client):
    body = wait(client, post(client, make_xlsx(ROWS)).json()["id"])
    hit = next(x for x in body["rows"] if x["overall_status"] == "ESCALATE_TO_COMPLIANCE")
    case = client.get(f"/api/applicants/{hit['applicant_id']}", headers=USER_HEADERS)
    assert case.status_code == 200 and case.json()["full_name"] == "Muhammad Ali Khan"
    mine = client.get("/api/applicants?mine=true", headers=USER_HEADERS).json()
    assert len(mine) == 3                                    # the three screenable rows
    pdf = client.get(f"/api/applicants/{hit['applicant_id']}/evidence", headers=USER_HEADERS)
    assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-"


def test_threshold_and_monitor_apply_to_every_row(client, storage):
    body = wait(client, post(client, make_xlsx(ROWS), threshold=97, monitor="true").json()["id"])
    assert body["threshold"] == 97 and body["monitor"] is True
    with storage.pool().connection() as conn:
        rows = conn.execute("SELECT threshold, monitored FROM applicants").fetchall()
    assert len(rows) == 3 and all(r["threshold"] == 97 and r["monitored"] for r in rows)


def test_header_aliases_numeric_cnic_and_dates(client, storage):
    from datetime import datetime
    data = make_xlsx([["Muhammad Ali Khan", datetime(1975, 3, 4), "Pakistan", 4210112345671, "Abdul", "KPK"]],
                     header=["Applicant Name", "DOB", "Citizenship", "NIC", "Father Name", "Province/Territory"])
    body = wait(client, post(client, data).json()["id"])
    assert body["rows"][0]["state"] == "screened"
    with storage.pool().connection() as conn:
        a = conn.execute("SELECT dob, nationality, cnic, father_name, province FROM applicants").fetchone()
    assert (a["dob"], a["nationality"], a["cnic"], a["father_name"], a["province"]) == \
        ("1975-03-04", "Pakistan", "4210112345671", "Abdul", "KPK")


def test_csv_with_semicolons_bom_and_blank_lines(client):
    csv = "\ufeffName;CNIC\r\n\r\nMuhammad Ali Khan;4210112345671\r\nSara Noor;\r\n".encode("utf-8")
    body = wait(client, post(client, csv, name="people.csv").json()["id"])
    assert [x["row"] for x in body["rows"]] == [3, 4]          # the row numbers of the file as a person sees it
    assert body["rows"][0]["overall_status"] == "ESCALATE_TO_COMPLIANCE"


def test_docx_table(client):
    body = wait(client, post(client, make_docx(ROWS), name="applicants.docx").json()["id"])
    assert body["total"] == 4 and body["counts"]["ESCALATE_TO_COMPLIANCE"] == 1


def test_xls_file(client):
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet("a")
    for c, h in enumerate(["Full name", "Nationality"]):
        ws.write(0, c, h)
    ws.write(1, 0, "Muhammad Ali Khan")
    ws.write(1, 1, "Pakistan")
    out = io.BytesIO()
    wb.save(out)
    body = wait(client, post(client, out.getvalue(), name="old.xls").json()["id"])
    assert body["counts"]["ESCALATE_TO_COMPLIANCE"] == 1


@pytest.mark.parametrize("name,data,status,code", [
    ("notes.pdf", b"%PDF-1.4", 415, "BATCH_FILE_TYPE"),
    ("a.xlsx", b"", 400, "BATCH_FILE_EMPTY"),
    ("a.xlsx", b"this is not a zip file", 422, "BATCH_FILE_UNREADABLE"),
    ("a.docx", b"this is not a zip file", 422, "BATCH_FILE_UNREADABLE"),
    ("a.xlsx", make_xlsx([["Sara Noor"]], header=["Surname", "Town"]), 422, "BATCH_NO_NAME_COLUMN"),
    ("a.xlsx", make_xlsx([]), 422, "BATCH_NO_ROWS"),
    ("a.xlsx", make_xlsx([["", "x"], ["\u0645\u062d\u0645\u062f \u0639\u0644\u06cc", ""]]), 422, "BATCH_NO_VALID_ROWS"),
    ("a.docx", b"PK\x05\x06" + b"\x00" * 18, 422, "BATCH_FILE_UNREADABLE"),
])
def test_unusable_files_are_refused_with_a_reason(client, name, data, status, code):
    r = post(client, data, name=name)
    assert r.status_code == status and r.json()["error"]["code"] == code
    assert client.get("/api/applicants?mine=true", headers=USER_HEADERS).json() == []


def test_docx_without_a_table(client):
    xml = ('<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           "<w:body><w:p><w:r><w:t>Just text</w:t></w:r></w:p></w:body></w:document>")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("word/document.xml", xml)
    r = post(client, out.getvalue(), name="a.docx")
    assert r.status_code == 422 and r.json()["error"]["code"] == "BATCH_NO_TABLE"


def test_too_many_rows_and_too_large_files(client, monkeypatch):
    from app import config
    monkeypatch.setattr(config, "BATCH_MAX_ROWS", 3)
    r = post(client, make_xlsx([["Person Number " + "x" * i] for i in range(4)]))
    assert r.status_code == 422 and r.json()["error"]["code"] == "BATCH_TOO_MANY_ROWS"
    monkeypatch.setattr(config, "BATCH_MAX_FILE_BYTES", 100)
    r = post(client, make_xlsx([["Sara Noor"]]))
    assert r.status_code == 413 and r.json()["error"]["code"] == "BATCH_FILE_TOO_LARGE"


def test_zip_that_expands_hugely_is_refused_before_it_is_unpacked():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"\0" * (90 * 1024 * 1024))
    with pytest.raises(batch_files.BatchFileError) as e:
        batch_files.parse_applicants(out.getvalue(), "bomb.docx", 500)
    assert e.value.code == "BATCH_FILE_TOO_LARGE"


def test_xml_entity_bombs_in_docx_are_refused():
    xml = ('<?xml version="1.0"?><!DOCTYPE d [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;">]>'
           '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>&b;</w:body></w:document>')
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("word/document.xml", xml)
    with pytest.raises(batch_files.BatchFileError):
        batch_files.parse_applicants(out.getvalue(), "evil.docx", 500)


# ---- access ------------------------------------------------------------------

def test_a_pending_user_cannot_start_a_batch(client):
    assert post(client, make_xlsx(ROWS), headers=PENDING_HEADERS).status_code == 403


def test_a_batch_is_private_to_the_person_who_uploaded_it_even_for_admins(client):
    bid = post(client, make_xlsx(ROWS)).json()["id"]
    wait(client, bid)
    for who in (USER2_HEADERS, API_HEADERS):
        for path in (f"/api/batches/{bid}", f"/api/batches/{bid}/results.xlsx", f"/api/batches/{bid}/evidence.zip"):
            assert client.get(path, headers=who).status_code == 404
        assert client.post(f"/api/batches/{bid}/cancel", headers=who).status_code == 404
    assert client.get("/api/batches/999", headers=USER_HEADERS).status_code == 404


# ---- failures, cancel, concurrency ---------------------------------------------

def test_one_failing_row_does_not_stop_the_others(client, monkeypatch):
    from app.screening import engine
    real = engine.screen

    def flaky(name, *a, **k):
        if name == "Sara Noor":
            raise RuntimeError("boom for Sara Noor")
        return real(name, *a, **k)

    monkeypatch.setattr(engine, "screen", flaky)
    body = wait(client, post(client, make_xlsx(ROWS)).json()["id"])
    rows = {x["row"]: x for x in body["rows"]}
    assert rows[5]["state"] == "failed" and rows[5]["overall_status"] is None
    assert "Sara" not in (rows[5]["error"] or "")            # the error never repeats what was in the row
    assert rows[2]["state"] == "screened" and rows[3]["state"] == "screened"
    assert body["status"] == "done" and body["counts"]["failed"] == 1


def _slow_screening(monkeypatch):
    from app import main
    gate, started = threading.Event(), threading.Event()
    real = main._screen_one

    def slow(req, user, request=None, batch_id=None):
        started.set()
        gate.wait(10)
        return real(req, user, request, batch_id)

    monkeypatch.setattr(main, "_screen_one", slow)
    return gate, started


def test_cancel_stops_after_the_current_row_and_unscreened_rows_stay_pending(client, monkeypatch):
    gate, started = _slow_screening(monkeypatch)
    bid = post(client, make_xlsx(ROWS)).json()["id"]
    assert started.wait(5)
    cancelled = client.post(f"/api/batches/{bid}/cancel", headers=USER_HEADERS)
    assert cancelled.status_code == 200
    gate.set()
    body = wait(client, bid)
    assert body["status"] == "cancelled"
    assert body["counts"]["screened"] == 1 and body["counts"]["pending"] == 2 and body["done"] == 2
    pending = [x for x in body["rows"] if x["state"] == "pending"]
    assert all(x["overall_status"] is None for x in pending)             # never shown as clear
    sheet = load_workbook(io.BytesIO(client.get(f"/api/batches/{bid}/results.xlsx", headers=USER_HEADERS).content))
    notes = [r[11].value for r in sheet["Results"].iter_rows(min_row=2)]
    assert sum("did not finish" in (n or "") for n in notes) == 2


def test_only_one_running_batch_per_person(client, monkeypatch):
    gate, started = _slow_screening(monkeypatch)
    first = post(client, make_xlsx(ROWS))
    assert started.wait(5)
    second = post(client, make_xlsx(ROWS))
    assert second.status_code == 409 and second.json()["error"]["code"] == "BATCH_ALREADY_RUNNING"
    assert post(client, make_xlsx(ROWS), headers=USER2_HEADERS).status_code == 202      # someone else may
    gate.set()
    wait(client, first.json()["id"])
    assert post(client, make_xlsx([["Sara Noor"]])).status_code == 202                   # and so may they, afterwards


def test_server_capacity_is_limited(client, monkeypatch):
    from app import config
    monkeypatch.setattr(config, "BATCH_MAX_RUNNING", 1)
    gate, started = _slow_screening(monkeypatch)
    first = post(client, make_xlsx(ROWS))
    assert started.wait(5)
    busy = post(client, make_xlsx(ROWS), headers=USER2_HEADERS)
    assert busy.status_code == 429 and busy.json()["error"]["code"] == "BATCH_BUSY"
    gate.set()
    wait(client, first.json()["id"])


def test_a_batch_cut_off_by_a_restart_is_marked_interrupted_not_clear(client, storage):
    gate = threading.Event()
    with storage.pool().connection() as conn:
        conn.execute("INSERT INTO batches (user_id, filename, threshold, total) VALUES (%s, 'x.xlsx', 85, 2)",
                     ("00000000-0000-4000-8000-0000000000b1",))
        conn.execute("INSERT INTO batch_rows (batch_id, row_no, full_name) VALUES (1, 2, 'Sara Noor'), (1, 3, 'Ali Raza')")
    assert storage.batch_interrupt_running() == 1
    body = client.get("/api/batches/1", headers=USER_HEADERS).json()
    assert body["status"] == "interrupted" and body["counts"]["pending"] == 2 and body["finished_at"]
    gate.set()


# ---- downloads and audit ----------------------------------------------------

def test_results_spreadsheet_has_every_row_and_never_a_formula(client):
    rows = ROWS + [['=HYPERLINK("http://evil.example","click")', "", "", "", "", ""],
                   ["+cmd|' /C calc'!A0", "", "", "", "", ""]]
    bid = post(client, make_xlsx(rows)).json()["id"]
    body = wait(client, bid)
    r = client.get(f"/api/batches/{bid}/results.xlsx", headers=USER_HEADERS)
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Results"]
    lines = list(ws.iter_rows(values_only=True))
    assert lines[0][:2] == ("Row", "Full name") and len(lines) == 1 + 6
    by_row = {l[0]: l for l in lines[1:]}
    assert by_row[2][7] == "Escalate to compliance" and by_row[2][10].startswith("CS-")
    assert by_row[4][7] == "Not screened" and "Full name" in by_row[4][11]
    for c in ws["B"][1:]:
        assert c.data_type != "f"                                       # never stored as a formula
    assert wb["Batch"]["B1"].value == "applicants.xlsx" and body["total"] == 6


def test_evidence_zip_has_the_pdfs_and_a_hash_manifest(client):
    bid = post(client, make_xlsx(ROWS)).json()["id"]
    wait(client, bid)
    r = client.get(f"/api/batches/{bid}/evidence.zip", headers=USER_HEADERS)
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(r.content))
    pdfs = [n for n in z.namelist() if n.endswith(".pdf")]
    assert len(pdfs) == 1 and pdfs[0].startswith("row-2-") and z.read(pdfs[0])[:5] == b"%PDF-"
    manifest = z.read("manifest.csv").decode()
    assert manifest.startswith("row,file,sha256") and pdfs[0] in manifest


def test_evidence_zip_is_404_when_nothing_was_found(client):
    bid = post(client, make_xlsx([["Completely Unrelated Person"]])).json()["id"]
    wait(client, bid)
    r = client.get(f"/api/batches/{bid}/evidence.zip", headers=USER_HEADERS)
    assert r.status_code == 404 and r.json()["error"]["code"] == "EVIDENCE_NOT_GENERATED"


def test_audit_trail_records_the_batch_without_any_names(client, storage):
    bid = post(client, make_xlsx(ROWS)).json()["id"]
    wait(client, bid)
    client.get(f"/api/batches/{bid}/results.xlsx", headers=USER_HEADERS)
    entries, _ = storage.audit_list(limit=100)
    actions = [e["action"] for e in entries]
    assert {"batch.start", "batch.finish", "batch.results.download", "screening.run"} <= set(actions)
    runs = [e for e in entries if e["action"] == "screening.run"]
    assert all(e["detail"]["batch_id"] == bid for e in runs)
    blob = " ".join(str(e) for e in entries)
    assert "Muhammad" not in blob and "Sara" not in blob


def test_a_name_in_urdu_is_flagged_on_its_row_and_never_looks_clear(client):
    rows = [["Sara Noor"], ["\u0645\u062d\u0645\u062f \u0639\u0644\u06cc"]]
    body = wait(client, post(client, make_xlsx(rows, header=["Full name"])).json()["id"])
    bad = body["rows"][1]
    assert bad["state"] == "invalid" and bad["overall_status"] is None and "Latin letters" in bad["error"]
    assert body["counts"]["AUTO_CLEAR"] == 1
