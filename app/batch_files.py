"""
Reading an uploaded file of applicants, and writing the results back out as a spreadsheet.

Supported inputs: Excel (.xlsx, .xls), CSV and a Word (.docx) table. The first row (or the first row that
names a "Full name" column, within the first few) holds the column names. Nothing here touches the network
or the database: the functions take bytes and return plain data, so they are easy to test.

Everything in an uploaded file is untrusted. Spreadsheets and Word files are zip archives, so their
declared uncompressed size is checked before anything is unpacked, and XML is parsed with defusedxml.
When the results are written back out, text is stored as text so a name like "=HYPERLINK(...)" can never
become a formula in the spreadsheet someone opens later.
"""

import csv
import io
import math
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime

from defusedxml import ElementTree as DefusedET

ALLOWED_EXTENSIONS = (".xlsx", ".xls", ".csv", ".docx")
_MAX_UNCOMPRESSED = 80 * 1024 * 1024       # what a zip based file may expand to
_HEADER_SEARCH_ROWS = 10                     # the column names must be in one of these rows

# the columns we use, and the headings people really write for them (compared in lowercase, letters and digits only)
_ALIASES = {
    "full_name": {"fullname", "name", "applicantname", "applicant", "customername", "clientname", "nameofapplicant"},
    "dob": {"dateofbirth", "dob", "birthdate", "dateofbirthyyyymmdd"},
    "nationality": {"nationality", "citizenship"},
    "cnic": {"cnic", "cnicnumber", "nic", "nicnumber", "nationalid", "idnumber", "idcardnumber"},
    "father_name": {"fatherorhusband", "fathername", "father", "husbandname", "husband", "fatherhusbandname",
                    "fatherorhusbandname", "fatherhusband"},
    "province": {"province", "state", "region", "provinceterritory"},
}
FIELDS = tuple(_ALIASES)


class BatchFileError(Exception):
    """The file cannot be used. `code` and `hint` become the API error."""

    def __init__(self, code: str, message: str, hint: str | None = None):
        super().__init__(message)
        self.code, self.message, self.hint = code, message, hint


@dataclass
class ParsedRow:
    row_no: int                      # as the person sees it: the spreadsheet row, or the table row in Word
    values: dict = field(default_factory=dict)


def extension_of(filename: str) -> str:
    m = re.search(r"(\.[A-Za-z0-9]+)$", filename or "")
    return m.group(1).lower() if m else ""


def _norm_header(text) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def _field_for(header) -> str | None:
    h = _norm_header(header)
    for name, aliases in _ALIASES.items():
        if h in aliases:
            return name
    return None


def _cell_text(value) -> str:
    """One cell as the text a person would read."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, datetime):
        return value.date().isoformat() if (value.hour, value.minute, value.second) == (0, 0, 0) else value.isoformat(" ", "seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""               # NaN or infinity (a damaged or odd cell): treat as empty instead of crashing
        # a CNIC or an id stored as a number: 4210112345671.0 must not become "4.210112345671e+12"
        return str(int(value)) if value == int(value) else repr(value)
    return re.sub(r"[ \t\r\n]+", " ", str(value)).strip()


def _check_zip(data: bytes) -> zipfile.ZipFile:
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
        total = sum(i.file_size for i in z.infolist())
    except (zipfile.BadZipFile, ValueError, OSError):
        raise BatchFileError("BATCH_FILE_UNREADABLE", "That file could not be read.",
                             "It may be damaged or not really an Excel or Word file. Open it, save it again, and retry.") from None
    if total > _MAX_UNCOMPRESSED:
        raise BatchFileError("BATCH_FILE_TOO_LARGE", "That file is too large once opened.",
                             "Split the applicants into smaller files.")
    return z


# --------------------------------------------------------------------------
# Readers: each returns a list of rows, each row a list of cell texts, plus the row numbers they came from
# --------------------------------------------------------------------------

def _read_xlsx(data: bytes, max_rows: int) -> list:
    _check_zip(data).close()
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise BatchFileError("BATCH_FILE_UNREADABLE", "That Excel file could not be read.",
                             "Open it in Excel, save it as .xlsx again, and retry.") from None
    try:
        ws = wb.worksheets[0] if wb.worksheets else None
        if ws is None:
            return []
        out = []
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            texts = [_cell_text(c) for c in row]
            if any(texts):
                out.append((i, texts))
            if len(out) > max_rows + _HEADER_SEARCH_ROWS + 1:
                break
        return out
    finally:
        wb.close()


def _read_xls(data: bytes, max_rows: int) -> list:
    try:
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        sheet = book.sheet_by_index(0)
    except Exception:
        raise BatchFileError("BATCH_FILE_UNREADABLE", "That Excel file could not be read.",
                             "Open it in Excel, save it as .xlsx, and retry.") from None
    out = []
    for r in range(sheet.nrows):
        texts = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    texts.append(_cell_text(xlrd.xldate_as_datetime(cell.value, book.datemode)))
                except Exception:
                    texts.append(_cell_text(cell.value))
            elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                texts.append("")
            else:
                texts.append(_cell_text(cell.value))
        if any(texts):
            out.append((r + 1, texts))
        if len(out) > max_rows + _HEADER_SEARCH_ROWS + 1:
            break
    return out


def _read_csv(data: bytes, max_rows: int) -> list:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("latin-1")
    sample = text[:4096]
    delimiter = max((",", ";", "\t", "|"), key=sample.count) if sample else ","
    out = []
    try:
        for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), start=1):
            texts = [_cell_text(c) for c in row]
            if any(texts):
                out.append((i, texts))
            if len(out) > max_rows + _HEADER_SEARCH_ROWS + 1:
                break
    except csv.Error:
        raise BatchFileError("BATCH_FILE_UNREADABLE", "That CSV file could not be read.",
                             "Save it again from Excel as CSV and retry.") from None
    return out


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _read_docx(data: bytes, max_rows: int) -> list:
    z = _check_zip(data)
    try:
        xml = z.read("word/document.xml")
    except KeyError:
        raise BatchFileError("BATCH_FILE_UNREADABLE", "That Word file could not be read.",
                             "Open it in Word, save it as .docx again, and retry.") from None
    finally:
        z.close()
    try:
        root = DefusedET.fromstring(xml)
    except Exception:
        raise BatchFileError("BATCH_FILE_UNREADABLE", "That Word file could not be read.",
                             "Open it in Word, save it as .docx again, and retry.") from None
    tables = list(root.iter(_W + "tbl"))
    if not tables:
        raise BatchFileError("BATCH_NO_TABLE", "That Word file has no table.",
                             "Put one applicant per row in a Word table, with the column names in its first row.")
    out = []
    for i, tr in enumerate(tables[0].iter(_W + "tr"), start=1):
        texts = []
        for tc in tr.findall(_W + "tc"):
            paragraphs = ["".join(t.text or "" for t in p.iter(_W + "t")) for p in tc.iter(_W + "p")]
            texts.append(_cell_text(" ".join(p for p in paragraphs if p)))
        if any(texts):
            out.append((i, texts))
        if len(out) > max_rows + _HEADER_SEARCH_ROWS + 1:
            break
    return out


_READERS = {".xlsx": _read_xlsx, ".xls": _read_xls, ".csv": _read_csv, ".docx": _read_docx}


def parse_applicants(data: bytes, filename: str, max_rows: int) -> list:
    """
    The applicants in an uploaded file, as ParsedRow(row_no, values) with values keyed by name, dob, ... (only
    the columns the file has; empty cells are ''). Blank rows are skipped. Raises BatchFileError if the file
    cannot be read, has no Full name column, has no applicants, or has more than `max_rows`.
    """
    ext = extension_of(filename)
    reader = _READERS.get(ext)
    if reader is None:
        raise BatchFileError("BATCH_FILE_TYPE", "That file type is not supported.",
                             "Use an Excel file (.xlsx or .xls), a CSV, or a Word file (.docx).")
    rows = reader(data, max_rows)

    header_at, columns = None, {}
    for pos, (_, texts) in enumerate(rows[:_HEADER_SEARCH_ROWS]):
        found = {}
        for idx, text in enumerate(texts):
            f = _field_for(text)
            if f and f not in found.values():
                found[idx] = f
        if "full_name" in found.values():
            header_at, columns = pos, found
            break
    if header_at is None:
        raise BatchFileError("BATCH_NO_NAME_COLUMN", "The file has no Full name column.",
                             "Put the column names in the first row, and call the applicant's name column Full name.")

    parsed = []
    for row_no, texts in rows[header_at + 1:]:
        values = {f: (texts[idx] if idx < len(texts) else "") for idx, f in columns.items()}
        if any(values.values()):
            parsed.append(ParsedRow(row_no, values))
    if not parsed:
        raise BatchFileError("BATCH_NO_ROWS", "The file has no applicants.",
                             "Add one applicant per row under the column names.")
    if len(parsed) > max_rows:
        raise BatchFileError("BATCH_TOO_MANY_ROWS", f"That file has more than {max_rows} applicants.",
                             f"Split it into files of at most {max_rows} rows.")
    return parsed


# --------------------------------------------------------------------------
# Writing the results
# --------------------------------------------------------------------------

_OUTCOME = {"ESCALATE_TO_COMPLIANCE": "Escalate to compliance", "MANUAL_REVIEW": "Needs manual review",
            "AUTO_CLEAR": "Clear"}
_STATE_NOTE = {"pending": "Not screened (the batch did not finish)"}
HEADERS = ["Row", "Full name", "Date of birth", "Nationality", "CNIC", "Father or husband", "Province", "Outcome",
           "Sanctions matches", "News articles", "PEP matches", "Case reference", "Note"]


def results_workbook(batch: dict, rows: list, case_ref) -> bytes:
    """
    The batch as an .xlsx a reviewer can filter. `rows` are database.batch_rows() rows, `case_ref` maps a row to
    its case reference. A row that was not screened says why in Note: it is never left looking clear.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        if r["state"] == "screened":
            outcome, note = _OUTCOME.get(r["overall_status"], r["overall_status"] or ""), ""
        elif r["state"] == "pending":
            outcome, note = "Not screened", _STATE_NOTE["pending"]
        else:
            outcome, note = "Not screened", r.get("error") or "This row could not be screened"
        ws.append([r["row_no"], r["full_name"], r.get("dob") or "", r.get("nationality") or "", r.get("cnic") or "",
                   r.get("father_name") or "", r.get("province") or "", outcome,
                   r["sanctions"] if r["state"] == "screened" else "", r["news"] if r["state"] == "screened" else "",
                   r["pep"] if r["state"] == "screened" else "",
                   case_ref(r) if r["state"] == "screened" else "", note])
    # text stays text: a name that starts with = + - or @ must not turn into a formula when the file is opened
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, str):
                c.data_type = "s"
    for i, width in enumerate([6, 32, 14, 16, 18, 28, 14, 24, 10, 10, 10, 22, 48], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Batch")
    for line in [("File", batch["filename"]), ("Threshold", batch["threshold"]), ("Rows", batch["total"]),
                 ("Status", batch["status"]), ("Started", batch["created_at"]), ("Finished", batch.get("finished_at") or "")]:
        info.append(list(line))
    for row in info.iter_rows():
        for c in row:
            if isinstance(c.value, str):
                c.data_type = "s"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
